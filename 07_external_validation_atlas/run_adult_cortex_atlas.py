from __future__ import annotations

import itertools
import json
import math
import os
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
from scipy.stats import hypergeom


ROOT = Path(
    os.environ.get(
        "EXTERNAL_VALIDATION_ROOT",
        str(Path(__file__).resolve().parent / "work"),
    )
)
WORK = ROOT / "adult_cortex_atlas"
SOURCE = WORK / "source"
OUT = WORK / "results"
OUT.mkdir(parents=True, exist_ok=True)

WEIGHTS = SOURCE / "final_pls1_gene_weights.csv"
LAMINAR_XLSX = SOURCE / "41467_2025_62793_MOESM4_ESM.xlsx"
REGIONAL_XLSX = SOURCE / "41467_2025_62793_MOESM5_ESM.xlsx"
ENSEMBL_MAPPING = SOURCE / "ensembl_gene_symbol_mapping.csv"
META_XLSX = SOURCE / "41467_2025_62793_MOESM3_ESM.xlsx"
FULL_LAMINAR = SOURCE / "MOESM8_laminar_markers"

LAYER_SHEETS = {
    "L1": "Marker gene - L1",
    "L2": "Marker gene -L2",
    "L3": "Marker gene -L3",
    "L4": "Marker gene -L4",
    "L5": "Marker gene -L5",
    "L6": "Marker gene -L6",
    "WM": "Marker gene -WM",
}
PRIMARY_LAYERS = ["L1", "L2", "L3", "L4", "L5", "L6"]
SET_ORDER = ["PLS1_positive", "PLS1_negative"]


def bh(pvalues: pd.Series | np.ndarray) -> np.ndarray:
    p = np.asarray(pvalues, dtype=float)
    out = np.full(len(p), np.nan)
    ok = np.isfinite(p)
    vals = p[ok]
    if len(vals) == 0:
        return out
    order = np.argsort(vals)
    ranked = vals[order]
    adjusted = ranked * len(ranked) / np.arange(1, len(ranked) + 1)
    adjusted = np.minimum.accumulate(adjusted[::-1])[::-1]
    adjusted = np.minimum(adjusted, 1.0)
    restored = np.empty_like(adjusted)
    restored[order] = adjusted
    out[np.where(ok)[0]] = restored
    return out


def exact_signflip_greater(values: np.ndarray) -> float:
    values = np.asarray(values, dtype=float)
    values = values[np.isfinite(values)]
    if len(values) == 0:
        return np.nan
    observed = values.mean()
    null = []
    for signs in itertools.product((-1.0, 1.0), repeat=len(values)):
        null.append(np.mean(values * np.asarray(signs)))
    null = np.asarray(null)
    return float((np.sum(null >= observed - 1e-15) + 1) / (len(null) + 1))


weights = pd.read_csv(WEIGHTS)
weights["gene"] = weights["gene"].astype(str).str.upper()
universe = set(weights["gene"])
gene_sets = {
    "PLS1_positive": set(weights.loc[weights["gene_set"].isin(["PLS1+", "PLS1_positive"]), "gene"]),
    "PLS1_negative": set(weights.loc[weights["gene_set"].isin(["PLS1-", "PLS1_negative"]), "gene"]),
}


def hypergeom_row(domain: str, set_name: str, markers: set[str], family: str) -> dict:
    markers = markers & universe
    genes = gene_sets[set_name] & universe
    overlap = markers & genes
    n_universe = len(universe)
    expected = len(markers) * len(genes) / n_universe
    fold = len(overlap) / expected if expected else np.nan
    p = hypergeom.sf(len(overlap) - 1, n_universe, len(genes), len(markers))
    return {
        "family": family,
        "domain": domain,
        "gene_set": set_name,
        "universe_N": n_universe,
        "set_N": len(genes),
        "marker_N_in_universe": len(markers),
        "overlap_N": len(overlap),
        "expected_overlap": expected,
        "enrichment_fold": fold,
        "p_one_sided": p,
        "overlap_genes": ";".join(sorted(overlap)),
    }


# Primary: frozen gene sets among the authors' published top laminar marker lists.
laminar_rows = []
for layer, sheet in LAYER_SHEETS.items():
    tab = pd.read_excel(LAMINAR_XLSX, sheet_name=sheet)
    gene_col = next(c for c in tab.columns if str(c).strip().lower() == "gene")
    markers = set(tab[gene_col].dropna().astype(str).str.upper())
    for set_name in SET_ORDER:
        fam = "primary_12" if layer in PRIMARY_LAYERS else "WM_boundary_control"
        laminar_rows.append(hypergeom_row(layer, set_name, markers, fam))

laminar = pd.DataFrame(laminar_rows)
primary_idx = laminar["family"].eq("primary_12")
laminar["q_global_primary_12"] = np.nan
laminar.loc[primary_idx, "q_global_primary_12"] = bh(laminar.loc[primary_idx, "p_one_sided"])
laminar["significant_primary_q05"] = laminar["q_global_primary_12"].lt(0.05)
laminar.to_csv(OUT / "laminar_top_marker_enrichment_globalBH.csv", index=False)


# Parse spatial sample metadata with merged-cell values carried downward.
wb = openpyxl.load_workbook(META_XLSX, read_only=False, data_only=True)
ws = wb["Sample information"]
meta_rows = []
current_experiment = None
current_donor = None
for r in range(7, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, 13)]
    if vals[0] is not None:
        current_experiment = str(vals[0])
    if vals[1] is not None:
        current_donor = str(vals[1])
    sample = vals[2]
    if sample is None:
        continue
    meta_rows.append(
        {
            "experiment": current_experiment,
            "donor": current_donor,
            "sample": str(sample),
            "region": "DLPFC" if vals[4] == "DLPFC (MFG)" else vals[4],
            "cortical_region": vals[3],
            "age": vals[8],
            "sex_ancestry": vals[9],
            "brodmann_area": vals[11],
        }
    )
wb.close()
meta = pd.DataFrame(meta_rows)
spatial_meta = meta[meta["experiment"].str.contains("Stereo", case=False, na=False)].copy()
spatial_meta.to_csv(OUT / "spatial_sample_metadata_parsed.csv", index=False)
coverage = (
    spatial_meta.groupby("region", dropna=False)
    .agg(n_slides=("sample", "nunique"), n_donors=("donor", "nunique"))
    .reset_index()
    .sort_values("region")
)
coverage.to_csv(OUT / "spatial_atlas_region_coverage.csv", index=False)
sample_to_donor = dict(zip(spatial_meta["sample"], spatial_meta["donor"]))
sample_to_region = dict(zip(spatial_meta["sample"], spatial_meta["region"]))


# Sensitivity: within each layer's full marker table, compare mean logFC of each
# frozen set with other AHBA-universe genes, per slide; aggregate to donor before
# exact sign-flip inference.
slide_rows = []
donor_rows = []
for layer in PRIMARY_LAYERS:
    tab = pd.read_csv(FULL_LAMINAR / f"{layer}.csv", low_memory=False)
    gene_col = tab.columns[0]
    tab[gene_col] = tab[gene_col].astype(str).str.upper()
    tab = tab[tab[gene_col].isin(universe)].drop_duplicates(gene_col).set_index(gene_col)
    fc_cols = [c for c in tab.columns if c.endswith("_avg_log2FC")]
    for set_name in SET_ORDER:
        in_set = tab.index.isin(gene_sets[set_name])
        for col in fc_cols:
            sample = col[: -len("_avg_log2FC")]
            vals = pd.to_numeric(tab[col], errors="coerce")
            a = vals[in_set & vals.notna()].to_numpy(float)
            b = vals[(~in_set) & vals.notna()].to_numpy(float)
            if len(a) < 5 or len(b) < 20:
                continue
            pooled = math.sqrt(((len(a) - 1) * np.var(a, ddof=1) + (len(b) - 1) * np.var(b, ddof=1)) / (len(a) + len(b) - 2))
            effect = (np.mean(a) - np.mean(b)) / pooled if pooled > 0 else np.nan
            slide_rows.append(
                {
                    "layer": layer,
                    "gene_set": set_name,
                    "sample": sample,
                    "donor": sample_to_donor.get(sample),
                    "region": sample_to_region.get(sample),
                    "n_set_genes": len(a),
                    "n_background_genes": len(b),
                    "standardized_effect": effect,
                }
            )

slide_effects = pd.DataFrame(slide_rows)
slide_effects.to_csv(OUT / "laminar_slide_level_standardized_effects.csv", index=False)
mapped = slide_effects.dropna(subset=["donor"]).copy()
donor_effects = (
    mapped.groupby(["layer", "gene_set", "donor"], as_index=False)
    .agg(donor_mean_effect=("standardized_effect", "mean"), n_slides=("sample", "nunique"))
)
donor_effects.to_csv(OUT / "laminar_donor_level_standardized_effects.csv", index=False)
for (layer, set_name), sub in donor_effects.groupby(["layer", "gene_set"]):
    vals = sub["donor_mean_effect"].to_numpy(float)
    donor_rows.append(
        {
            "layer": layer,
            "gene_set": set_name,
            "n_donors": len(vals),
            "mean_donor_effect": np.mean(vals),
            "median_donor_effect": np.median(vals),
            "donors_positive": int(np.sum(vals > 0)),
            "p_exact_signflip_greater": exact_signflip_greater(vals),
        }
    )
donor_summary = pd.DataFrame(donor_rows)
donor_summary["q_global_12"] = bh(donor_summary["p_exact_signflip_greater"])
donor_summary["significant_q05"] = donor_summary["q_global_12"].lt(0.05)
donor_summary.to_csv(OUT / "laminar_cross_donor_sensitivity_globalBH.csv", index=False)


# Exploratory regional heterogeneity within each excitatory subclass.
regional = pd.read_excel(REGIONAL_XLSX, sheet_name="Marker gene-subclass-region-EXC")
mapping = pd.read_csv(ENSEMBL_MAPPING)
mapping["ensembl_gene_id"] = mapping["ensembl_gene_id"].astype(str).str.split(".").str[0]
mapping_lookup = dict(zip(mapping["ensembl_gene_id"], mapping["gene_symbol"]))
regional["source_gene_id"] = regional["gene"].astype(str)
regional["gene"] = regional["source_gene_id"].str.split(".").str[0].map(mapping_lookup)
regional["gene"] = regional["gene"].astype("string").str.upper()
regional["avg_log2FC"] = pd.to_numeric(regional["avg_log2FC"], errors="coerce")
regional["p_val_adj"] = pd.to_numeric(regional["p_val_adj"], errors="coerce")
regional_rows = []
for (subclass, region), sub in regional.groupby(["subclass", "region"], dropna=False):
    markers = set(sub.loc[(sub["p_val_adj"] < 0.05) & (sub["avg_log2FC"] > 0), "gene"])
    for set_name in SET_ORDER:
        row = hypergeom_row(f"{subclass}__{region}", set_name, markers, "regional_exploratory")
        row["subclass"] = subclass
        row["region"] = region
        regional_rows.append(row)
regional_enrichment = pd.DataFrame(regional_rows)
regional_enrichment["q_global_all_regional_tests"] = bh(regional_enrichment["p_one_sided"])
regional_enrichment["significant_q05"] = regional_enrichment["q_global_all_regional_tests"].lt(0.05)
regional_enrichment.to_csv(OUT / "regional_excitatory_subclass_enrichment_globalBH.csv", index=False)

regional_summary = (
    regional_enrichment.groupby(["subclass", "gene_set"], as_index=False)
    .agg(
        n_regions_tested=("region", "nunique"),
        n_regions_q05=("significant_q05", "sum"),
        min_global_q=("q_global_all_regional_tests", "min"),
        max_enrichment_fold=("enrichment_fold", "max"),
    )
)
regional_summary.to_csv(OUT / "regional_excitatory_subclass_summary.csv", index=False)


# Compact review figure.
plt.rcParams.update({"font.family": "Arial", "font.size": 14})
fig = plt.figure(figsize=(12.2, 6.8), constrained_layout=True)
gs = fig.add_gridspec(2, 2, width_ratios=[1.05, 1.55], height_ratios=[1.0, 0.72])

ax1 = fig.add_subplot(gs[0, 0])
heat = laminar[laminar["family"].eq("primary_12")].pivot(index="gene_set", columns="domain", values="enrichment_fold").reindex(index=SET_ORDER, columns=PRIMARY_LAYERS)
im = ax1.imshow(np.log2(heat.clip(lower=2**-1.5)), cmap="RdBu_r", vmin=-1.5, vmax=1.5, aspect="auto")
for i, set_name in enumerate(SET_ORDER):
    for j, layer in enumerate(PRIMARY_LAYERS):
        rr = laminar[(laminar.domain == layer) & (laminar.gene_set == set_name)].iloc[0]
        star = "*" if rr["q_global_primary_12"] < 0.05 else ""
        ax1.text(j, i, f"{rr.enrichment_fold:.1f}×{star}\n({int(rr.overlap_N)})", ha="center", va="center", fontsize=13)
ax1.set_xticks(range(len(PRIMARY_LAYERS)), PRIMARY_LAYERS)
ax1.set_yticks(range(len(SET_ORDER)), ["PLS1+", "PLS1−"])
ax1.set_title("A  Frozen PLS sets in adult laminar markers\nfold enrichment; * global BH q<0.05")
fig.colorbar(im, ax=ax1, fraction=0.045, label="log2(fold enrichment)")

ax2 = fig.add_subplot(gs[0, 1])
subclasses = sorted(regional_summary["subclass"].dropna().unique())
reg_heat = regional_summary.pivot(index="gene_set", columns="subclass", values="n_regions_q05").reindex(index=SET_ORDER, columns=subclasses).fillna(0)
im2 = ax2.imshow(reg_heat, cmap="Blues", vmin=0, vmax=max(1, reg_heat.to_numpy().max()), aspect="auto")
for i in range(reg_heat.shape[0]):
    for j in range(reg_heat.shape[1]):
        ax2.text(j, i, str(int(reg_heat.iloc[i, j])), ha="center", va="center", fontsize=13)
ax2.set_xticks(range(len(subclasses)), subclasses, rotation=55, ha="right")
ax2.set_yticks(range(len(SET_ORDER)), ["PLS1+", "PLS1−"])
ax2.set_title("B  Regional heterogeneity within excitatory subclasses\nnumber of regions at global BH q<0.05")
fig.colorbar(im2, ax=ax2, fraction=0.025, label="significant regions")

ax3 = fig.add_subplot(gs[1, :])
coverage_plot = coverage.sort_values(["n_slides", "region"], ascending=[False, True])
x = np.arange(len(coverage_plot))
ax3.bar(x, coverage_plot["n_slides"], color="#4678A6", label="Stereo-seq slides")
ax3.plot(x, coverage_plot["n_donors"], color="#B24C45", marker="o", linewidth=1.5, label="donors")
ax3.set_xticks(x, coverage_plot["region"], rotation=45, ha="right")
ax3.set_ylabel("count")
ax3.set_title("C  Adult atlas sampling across 14 cortical regions")
ax3.legend(frameon=False, ncol=2)
ax3.spines[["top", "right"]].set_visible(False)

fig.suptitle("Independent neurotypical adult multicortical context for the frozen AHBA PLS signature", fontsize=13, fontweight="bold")
fig.savefig(OUT / "adult_cortex_atlas_adult_cortex_validation.png", dpi=600, bbox_inches="tight")
fig.savefig(OUT / "adult_cortex_atlas_adult_cortex_validation.pdf", dpi=600, bbox_inches="tight")
plt.close(fig)


primary_sig = laminar[primary_idx & laminar["significant_primary_q05"]]
acceptance = {
    "stage": 6,
    "status": "PASS",
    "primary_tests": int(primary_idx.sum()),
    "primary_significant_q05": int(len(primary_sig)),
    "primary_significant_domains": primary_sig[["domain", "gene_set"]].to_dict("records"),
    "spatial_slides_parsed": int(spatial_meta["sample"].nunique()),
    "spatial_donors_parsed": int(spatial_meta["donor"].nunique()),
    "spatial_regions_parsed": int(spatial_meta["region"].nunique()),
    "regional_exploratory_tests": int(len(regional_enrichment)),
    "regional_exploratory_significant_q05": int(regional_enrichment["significant_q05"].sum()),
    "inference_boundary": "neurotypical laminar/regional context only; not an HD effect and not literal whole-cortex coverage",
}
(OUT / "adult_cortex_acceptance.json").write_text(json.dumps(acceptance, indent=2), encoding="utf-8")

print(json.dumps(acceptance, indent=2))
print("\nPrimary laminar results:")
print(laminar.loc[primary_idx, ["domain", "gene_set", "marker_N_in_universe", "overlap_N", "enrichment_fold", "p_one_sided", "q_global_primary_12"]].to_string(index=False))
print("\nCross-donor sensitivity:")
print(donor_summary.to_string(index=False))
print("\nRegional subclass summary:")
print(regional_summary.to_string(index=False))

# =============================================================================
# SCRIPT DOCUMENTATION
# Purpose: Test the reference transcriptional signature against an adult cortical laminar/cell-type atlas.
# Input source/location: Adult cortex atlas expression/marker data, HGNC/Ensembl mappings, and reference PLS gene sets from configurable inputs.
# Output location: Laminar/cell-type enrichment tables, global multiple-testing results, and supporting summaries in the external-validation results directory.
# Input/output notes: Inputs are de-identified/public analysis resources; outputs contain derived analysis products only.
# Main steps: Map atlas genes; construct layer/cell-type marker sets; test enrichment/directionality for reference gene sets; apply global correction; export results.
# Log location: No dedicated log file unless configured externally; progress is written to standard output.
# =============================================================================
